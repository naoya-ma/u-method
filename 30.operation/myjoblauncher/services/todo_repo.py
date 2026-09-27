import os
import tomllib
from dataclasses import dataclass
from pathlib import Path
from typing import List

import tomli_w
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from services.path_utils import expand_path

TODO_PATH = "config/todos.toml"

TARGET_OBSIDIAN = "Obsidian"
TARGET_EXCEL = "Excel"


# ========================
# ✅ データモデル
# ========================
@dataclass
class TodoItem:
    subject: str = ""            # 件名（一覧の主表示）
    content: str = ""            # 内容（Markdown本文。記入テンプレートで見出しを挿入できる）
    target: str = TARGET_OBSIDIAN  # 書き込み先（"Obsidian" | "Excel"）
    todo_id: str = ""            # 自動採番される安定識別子（例: "T001"）
    created_at: str = ""         # 作成日（YYYY-MM-DD）
    due_date: str = ""           # 期限（YYYY-MM-DD）。省略可
    done: bool = False
    done_at: str = ""            # 完了日（YYYY-MM-DD）


@dataclass
class TodoSettings:
    excel_path: str = ""         # Excel書き込み先ファイル（.xlsx）


@dataclass
class TodoTemplate:
    name: str                    # テンプレート名（例: "タスク"）
    content: str = ""            # 内容欄へ挿入するMarkdown見出し構成


# ✅ 初回起動時に [[templates]] が存在しない場合の初期登録セット
DEFAULT_TEMPLATES = [
    TodoTemplate(
        name="タスク",
        content="### 目的\n\n### 背景・経緯\n\n### GOAL\n\n### 進め方\n\n### クローズ条件\n\n### インプット\n\n### アウトプット\n",
    ),
    TodoTemplate(
        name="課題",
        content="### クローズ条件\n\n### 進め方\n\n### インプット\n\n### アウトプット\n",
    ),
    TodoTemplate(
        name="トラブル",
        content="### いつ\n\n### どこで\n\n### 何が起こったか\n\n### 影響\n\n### 暫定対策\n\n### 恒久対象\n",
    ),
]


# ========================
# ✅ 読み込み・書き込み（[[items]] と [settings] は互いを消さないようマージして書く）
# ========================
def _load_raw(path: str) -> dict:
    if not os.path.exists(path):
        return {}
    with open(path, "rb") as f:
        return tomllib.load(f)


def load_todos(path: str = TODO_PATH) -> List[TodoItem]:
    data = _load_raw(path)
    items = []
    for t in data.get("items", []):
        items.append(
            TodoItem(
                # "subject"は旧"text"（本文のみだった時代のキー）からのフォールバックに対応する
                subject=t.get("subject", t.get("text", "")),
                content=t.get("content", ""),
                target=t.get("target", TARGET_OBSIDIAN),
                todo_id=t.get("todo_id", ""),
                created_at=t.get("created_at", ""),
                due_date=t.get("due_date", ""),
                done=bool(t.get("done", False)),
                done_at=t.get("done_at", ""),
            )
        )

    assigned = False
    for item in items:
        if not item.todo_id:
            item.todo_id = next_todo_id(items)
            assigned = True
    if assigned:
        save_todos(items, path)

    return items


def save_todos(todos: List[TodoItem], path: str = TODO_PATH) -> None:
    """[settings] セクションは既存の内容のまま、[[items]] だけを書き換える"""
    data = _load_raw(path)
    data["items"] = [
        {
            "todo_id": t.todo_id,
            "subject": t.subject,
            "content": t.content,
            "target": t.target,
            "created_at": t.created_at,
            "due_date": t.due_date,
            "done": t.done,
            "done_at": t.done_at,
        }
        for t in todos
    ]
    with open(path, "wb") as f:
        tomli_w.dump(data, f)


def next_todo_id(items: List[TodoItem]) -> str:
    """次に採番するTODO ID（例: T001）。既存の最大値の次から採番し、欠番は再利用しない"""
    existing_numbers = [
        int(item.todo_id[1:])
        for item in items
        if item.todo_id.startswith("T") and item.todo_id[1:].isdigit()
    ]
    return f"T{max(existing_numbers, default=0) + 1:03d}"


def load_todo_settings(path: str = TODO_PATH) -> TodoSettings:
    data = _load_raw(path)
    settings = data.get("settings", {})
    return TodoSettings(excel_path=settings.get("excel_path", ""))


def save_todo_settings(settings: TodoSettings, path: str = TODO_PATH) -> None:
    """[[items]] セクションは既存の内容のまま、[settings] だけを書き換える"""
    data = _load_raw(path)
    data["settings"] = {"excel_path": settings.excel_path}
    with open(path, "wb") as f:
        tomli_w.dump(data, f)


def load_templates(path: str = TODO_PATH) -> List[TodoTemplate]:
    """
    [[templates]] が一度も保存されていない（キー自体が無い）場合のみ、DEFAULT_TEMPLATESで初期登録する。
    ユーザーが全件削除した後の「空リスト」とは区別する（空リストならそのまま空で返す）
    """
    data = _load_raw(path)
    if "templates" not in data:
        templates = list(DEFAULT_TEMPLATES)
        save_templates(templates, path)
        return templates

    return [TodoTemplate(name=t.get("name", ""), content=t.get("content", "")) for t in data["templates"]]


def save_templates(templates: List[TodoTemplate], path: str = TODO_PATH) -> None:
    """[[items]]/[settings] セクションは既存の内容のまま、[[templates]] だけを書き換える"""
    data = _load_raw(path)
    data["templates"] = [{"name": t.name, "content": t.content} for t in templates]
    with open(path, "wb") as f:
        tomli_w.dump(data, f)


# ========================
# ✅ 書き込み先（Obsidian／Excel）への実書き込み
# ========================
def write_todo_to_obsidian(item: TodoItem, vault_dir: str, person_name: str) -> Path:
    """
    <vault_dir>/<person_name>/TODO.md へ完了済みTODOを見出しブロックとして追記する（無ければ新規作成）。
    日報本体（obsidian_repo.write_daily_report、日付単位で上書き）とは異なり、
    こちらは追記のみの「持ち越しTODO」専用の運用ノート。
    内容（content）は記入テンプレート由来の`### 見出し`を複数含みうるため、
    1行のリスト項目には押し込めず独立した見出しブロックにする
    """
    path = Path(expand_path(vault_dir)) / person_name / "TODO.md"
    path.parent.mkdir(parents=True, exist_ok=True)

    if not path.exists():
        path.write_text("# TODO\n\n", encoding="utf-8")

    with open(path, "a", encoding="utf-8") as f:
        f.write(f"## ✅ {item.subject}（完了: {item.done_at}）\n\n{item.content}\n\n---\n\n")

    return path


def write_todo_to_excel(item: TodoItem, excel_path: str) -> None:
    """excel_path の "TODO" シートへ完了済みTODOを1行追記する（ファイル・シートが無ければ新規作成）"""
    path = Path(excel_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        wb = load_workbook(path)
        is_new_sheet = "TODO" not in wb.sheetnames
        ws = wb.create_sheet("TODO") if is_new_sheet else wb["TODO"]
    else:
        wb = Workbook()
        wb.active.title = "TODO"
        ws = wb.active
        is_new_sheet = True

    if is_new_sheet:
        ws.append(["完了日", "件名", "内容", "期限"])
    ws.append([item.done_at, item.subject, item.content, item.due_date])
    ws.cell(row=ws.max_row, column=3).alignment = Alignment(wrap_text=True)

    wb.save(path)


# ========================
# ✅ 日報Markdownへの埋め込み用整形
# ========================
def build_todo_markdown(todos: List[TodoItem]) -> str:
    """
    その時点の全TODOをチェックリスト形式で整形する（完了済みは取り消し線）。
    一覧表示と同様、件名（＋期限）のみの簡潔な形式にする（内容の全文はここには含めない）
    """
    if not todos:
        return ""

    lines = []
    for t in todos:
        if t.done:
            lines.append(f"- [x] ~~{t.subject}~~（完了: {t.done_at}）")
        elif t.due_date:
            lines.append(f"- [ ] {t.subject}（期限: {t.due_date}）")
        else:
            lines.append(f"- [ ] {t.subject}")
    return "\n".join(lines)
