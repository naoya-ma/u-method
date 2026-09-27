"""
統合業務管理フェーズ4: ジョブの一括インポート／エクスポート（DESIGN.md「11.3」フェーズ4、
11.2 B3参照）。Excel(.xlsx)を、ジョブ追加・編集ダイアログ（`gui/job_editor.py`）と同じ
項目の固定日本語ヘッダーで読み書きする（エクスポートしたファイルをそのまま編集して
再インポートできる、往復可能な形式）。

スプレッドシート側には安定したユニークキーが無いのが普通のため、2つのインポートモードに
分ける（DESIGN.md 11.2 B3）:
- "add"（新規追加のみ）: 全行を新規ジョブとして追加する。job_id列は無視する
- "update"（job_id列明示による更新）: job_id列必須。既存ジョブと一致する行だけ内容を
  更新する。一致しない場合はエラーとして報告するのみで、新規追加はしない（誤って
  無関係な行を新規ジョブとして紛れ込ませないよう、安全側に倒す設計）

このモジュールはメモリ上のMenuItemを組み立てる（`import_jobs_from_excel`）／実際に
既存一覧へ適用する（`apply_import_result`）ところまでを担当し、ファイルへの永続化
（`persist_jobs`/`save_config`）は呼び出し側（GUI）が行う。呼び出し側は永続化の前に
`ImportResult`の内容（追加・更新・エラー件数）を必ず確認画面で提示すること
（誤ったファイルを一括インポートしてしまう事故を防ぐため）
"""

import os
from dataclasses import dataclass, replace

import openpyxl

from services.config_loader import (
    JOB_TYPE_FIXED,
    JOB_TYPE_FLEXIBLE,
    JOB_TYPES,
    MenuItem,
    next_job_id,
)
from services import tz_repo

COL_JOB_ID = "job_id"
COL_NAME = "ジョブ名"
COL_JOB_TYPE = "種別"
COL_TIER = "Tier"
COL_CATEGORY = "カテゴリ"
COL_COMMAND = "コマンド"
COL_IS_ASYNC = "非同期実行"
COL_SCHEDULE = "スケジュール"
COL_MANUAL = "業務マニュアル"
COL_MATURITY = "成熟度"
COL_NOTE = "特記事項"
COL_CREATED_AT = "作成日"
COL_UPDATED_AT = "最終更新日"
COL_ENABLED = "有効"

# ✅ エクスポート列の並び順。job_id/作成日/最終更新日/有効は参考情報も兼ねて含めるが、
# インポート時はjob_id（updateモードのみ必須）以外読み取らない（自動管理のため）
EXPORT_COLUMNS = [
    COL_JOB_ID,
    COL_NAME,
    COL_JOB_TYPE,
    COL_TIER,
    COL_CATEGORY,
    COL_COMMAND,
    COL_IS_ASYNC,
    COL_SCHEDULE,
    COL_MANUAL,
    COL_MATURITY,
    COL_NOTE,
    COL_CREATED_AT,
    COL_UPDATED_AT,
    COL_ENABLED,
]

MODE_ADD = "add"
MODE_UPDATE = "update"

_TRUE_TOKENS = {"true", "1", "有", "はい", "○", "yes", "y", "on"}


def export_jobs_to_excel(menu_items: list[MenuItem], path: str) -> None:
    """ジョブ一覧をExcel(.xlsx)へ出力する（そのまま編集して一括インポートできる往復可能な形式）"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "jobs"
    ws.append(EXPORT_COLUMNS)

    for item in menu_items:
        ws.append(
            [
                item.job_id,
                item.name,
                item.job_type,
                item.tier,
                item.category,
                item.command,
                "TRUE" if item.is_async else "FALSE",
                item.schedule,
                item.manual,
                item.maturity or "",
                item.note,
                item.created_at,
                item.updated_at,
                "TRUE" if item.enabled else "FALSE",
            ]
        )

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


@dataclass
class ImportRowError:
    row_number: int  # 1始まり（ヘッダー行=1、データ行は2から）
    message: str


@dataclass
class ImportResult:
    mode: str
    added: list[MenuItem]
    # (更新対象の既存job_id, 更新後の内容を持つ仮のMenuItem)。適用は apply_import_result() で行う
    updates: list[tuple[str, MenuItem]]
    errors: list[ImportRowError]


def _parse_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in _TRUE_TOKENS


def _parse_maturity(value) -> int:
    raw = str(value or "").strip()
    if not raw:
        return 0
    try:
        n = int(float(raw))
    except ValueError:
        return 0
    return n if 0 <= n <= 5 else 0


def _read_rows(path: str) -> tuple[list[str], list[tuple[int, dict]]]:
    """
    ヘッダー行から列名→セル値の対応付けを行い、以降の行を(元のExcel上の行番号, {ヘッダー名: セル値})の
    タプルとして返す。列の並び順は問わない（ヘッダー名の完全一致のみで対応付けるため、手作業で
    列を入れ替えても問題ない）。値が全て空のデータ行はスキップするが、**スキップした行があっても
    後続の行番号はExcel上の実際の行番号のまま**にする（スキップ分を詰めて数え直すと、エラー報告の
    行番号が実際のファイルの行とずれてしまい、修正時にユーザーが誤った行を探してしまうため）
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    data_rows = []
    for offset, raw_row in enumerate(rows[1:]):
        row_number = offset + 2  # ヘッダーが1行目、データは2行目から
        if raw_row is None or all(c is None for c in raw_row):
            continue
        row = {header[i]: raw_row[i] for i in range(len(header)) if i < len(raw_row)}
        data_rows.append((row_number, row))

    return header, data_rows


def import_jobs_from_excel(path: str, mode: str, existing_items: list[MenuItem]) -> ImportResult:
    """
    Excel(.xlsx)を読み取り、追加・更新の候補（ImportResult）を組み立てる。
    既存の`existing_items`は一切変更しない（副作用なし。確認後の適用は`apply_import_result()`で行う）
    """
    if mode not in (MODE_ADD, MODE_UPDATE):
        raise ValueError(f"不明なインポートモード: {mode}")

    header, rows = _read_rows(path)

    missing = [c for c in (COL_NAME,) if c not in header]
    if mode == MODE_UPDATE and COL_JOB_ID not in header:
        missing.append(COL_JOB_ID)
    if missing:
        raise ValueError(f"必須の列が見つかりません: {', '.join(missing)}")

    existing_by_id = {item.job_id: item for item in existing_items if item.job_id}
    # ✅ 追加モードでの新規採番は、同じ一括インポート内で増えていく件数も踏まえて連番が
    # ずれないよう、採番済みの項目を都度working_itemsへ積み上げてnext_job_id()へ渡す
    working_items = list(existing_items)

    added: list[MenuItem] = []
    updates: list[tuple[str, MenuItem]] = []
    errors: list[ImportRowError] = []
    today = tz_repo.today().isoformat()

    for row_number, row in rows:
        name = str(row.get(COL_NAME) or "").strip()
        if not name:
            errors.append(ImportRowError(row_number, "ジョブ名が空のためスキップしました"))
            continue

        job_type_raw = str(row.get(COL_JOB_TYPE) or "").strip()
        job_type = job_type_raw if job_type_raw in JOB_TYPES else JOB_TYPE_FIXED
        maturity = _parse_maturity(row.get(COL_MATURITY))
        command = str(row.get(COL_COMMAND) or "").strip()
        is_skeleton = maturity == 0

        if not command and job_type != JOB_TYPE_FLEXIBLE and not is_skeleton:
            errors.append(
                ImportRowError(
                    row_number,
                    f"「{name}」: コマンドが未設定です"
                    "（非定型ジョブ・成熟度未設定〈骨組み〉以外は必須です）",
                )
            )
            continue

        tier = str(row.get(COL_TIER) or "").strip()
        category = str(row.get(COL_CATEGORY) or "").strip()
        is_async = _parse_bool(row.get(COL_IS_ASYNC))
        schedule = str(row.get(COL_SCHEDULE) or "").strip()
        manual = str(row.get(COL_MANUAL) or "").strip()
        note = str(row.get(COL_NOTE) or "").strip()

        if mode == MODE_ADD:
            new_item = MenuItem(
                name=name,
                command=command,
                is_async=is_async,
                schedule=schedule,
                note=note,
                manual=manual,
                category=category,
                job_type=job_type,
                tier=tier,
                job_id=next_job_id(working_items),
                created_at=today,
                updated_at=today,
                maturity=maturity,
            )
            working_items.append(new_item)
            added.append(new_item)
            continue

        # mode == MODE_UPDATE
        job_id = str(row.get(COL_JOB_ID) or "").strip()
        if not job_id:
            errors.append(ImportRowError(row_number, f"「{name}」: job_idが空のためスキップしました"))
            continue

        target = existing_by_id.get(job_id)
        if target is None:
            errors.append(ImportRowError(row_number, f"「{name}」: job_id「{job_id}」が見つかりません"))
            continue

        # 更新後の内容を持つ「仮の」MenuItem（job_id/作成日/有効フラグは既存のものを維持する）
        updated_item = replace(
            target,
            name=name,
            command=command,
            is_async=is_async,
            schedule=schedule,
            note=note,
            manual=manual,
            category=category,
            job_type=job_type,
            tier=tier,
            maturity=maturity,
            updated_at=today,
        )
        updates.append((job_id, updated_item))

    return ImportResult(mode=mode, added=added, updates=updates, errors=errors)


def apply_import_result(existing_items: list[MenuItem], result: ImportResult) -> list[MenuItem]:
    """
    確認後に実際へ適用する。`existing_items`は変更せず、適用結果の新しいリストを返す
    （呼び出し側で`menu_items[:] = apply_import_result(...)`のように差し替えて使う）
    """
    updates_by_id = {job_id: item for job_id, item in result.updates}
    merged = [updates_by_id.get(item.job_id, item) for item in existing_items]
    merged.extend(result.added)
    return merged
